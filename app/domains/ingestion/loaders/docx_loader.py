import io

from docx import Document

from app.domains.ingestion.loaders.base import IDocumentLoader, LoadedDocument, SourceInput
from app.shared.exceptions import IngestionError


class DOCXLoader(IDocumentLoader):
    kind = "docx"

    async def load(self, source: SourceInput) -> list[LoadedDocument]:
        if not source.content:
            raise IngestionError("DOCX content is empty")
        try:
            doc = Document(io.BytesIO(source.content))
        except Exception as exc:
            raise IngestionError(f"Failed to open DOCX: {exc}") from exc
        text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        if not text:
            raise IngestionError("DOCX contains no text")
        return [LoadedDocument(text=text, metadata={"filename": source.filename})]