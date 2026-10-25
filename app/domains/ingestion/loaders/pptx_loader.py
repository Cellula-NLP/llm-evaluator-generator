import io

from pptx import Presentation

from app.domains.ingestion.loaders.base import IDocumentLoader, LoadedDocument, SourceInput
from app.shared.exceptions import IngestionError


class PPTXLoader(IDocumentLoader):
    kind = "pptx"

    async def load(self, source: SourceInput) -> list[LoadedDocument]:
        if not source.content:
            raise IngestionError("PPTX content is empty")
        try:
            prs = Presentation(io.BytesIO(source.content))
        except Exception as exc:
            raise IngestionError(f"Failed to open PPTX: {exc}") from exc

        docs: list[LoadedDocument] = []
        for i, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    parts.append(f"[Notes] {note}")
            if parts:
                docs.append(LoadedDocument(text="\n".join(parts), metadata={"slide": i, "filename": source.filename}))
        if not docs:
            raise IngestionError("PPTX contains no text")
        return docs